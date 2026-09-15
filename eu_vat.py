"""AB üye devletlerinin KDV oranları — tohum veri + TEDB'den otomatik yükseltme.

Neden tohum? AB'nin resmî referansı **TEDB** (Taxes in Europe Database) ve KDV oranları
sayfası bunu açıkça söylüyor. Ama TEDB bu tarihte başsız (headless) sorguya kapalı;
canlı olarak ölçüldü:

* ``GET /rest-api/configurations`` → 200, 29 ülke (çalışıyor, eşitlemede kullanılıyor).
* ``POST /rest-api/vatSearch`` → **boş gövde dâhil her istekte 500** (``TEDB-ERR-1``).
* ``GET /rest-api/vatSearch/export`` → 200 ve doğru sütun yapısında bir .xlsx veriyor,
  ama **veri satırı sıfır**: dışa aktarım sunucu oturumundaki son aramayı yazdığı için
  anonim istekte boş şablon dönüyor. Kategori, CN kodu ve geçmiş tarih denemeleri de
  boş döndü.

Bu yüzden veri elle derlenmiş bir tohumdan gelir. ``sync()`` her çalıştığında TEDB'yi
yoklar; TEDB veri döndürmeye başladığı gün tohum kendiliğinden resmî veriyle değişir. Bu,
``vat_lists.py``'nin Türk KDV listesinde kullandığı desenin aynısıdır.

**İki ayrı kanıt düzeyi vardır ve birbirinin yerine geçmezler:**

* ``verified`` — satır bir resmî anlık görüntüden **makine tarafından okundu** (kaynak URL +
  tarih + sha256). TEDB başsız çalışmadığı için bugün hiçbir satır bu düzeyde değil.
* ``expert_confirmed`` — tabloyu uygulamanın sahibi olan gümrük müşaviri **teyit etti**.
  İnsan teyidi makine okumasının yerine geçmez, ama "doğrulanmamış tohum" uyarısını da
  gerçeğe aykırı kılar. Teyit süresiz değildir: ``EU_VAT_CONFIRMATION_MAX_AGE_DAYS``
  (varsayılan 180 gün) geçince düşer, uyarı geri gelir ve oranı sık değişen ülkeler
  (``volatile``) yeniden öncelikli doğrulama listesine girer.

**İndirimli oran bir öneridir, tespit değildir.** Kullanıcı indirimli oranın GTİP faslından
otomatik seçilmesini istedi; uygulanıyor, ama AB Ek-III kategorileri ile GTİP faslı birebir
örtüşmez ve her üye devlet Ek-III'ü farklı uygular. Bu yüzden fasıl kuralı tetiklendiğinde
sonuç **her zaman** ``ambiguous=True`` döner, standart oran da adaylar arasında kalır ve
çağıranın bu değeri "doğrulandı" olarak göstermesi yasaktır (bkz. ``export_requirements``).
"""

from __future__ import annotations

import asyncio
import hashlib
import io
import json
import logging
import os
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import urlencode

import httpx

from security_firewall import SecurityViolation, validate_outbound_url
from trade_measures import USER_AGENT, official_ssl_context

logger = logging.getLogger(__name__)

DATA_FILE = "eu_vat_rates.json"
SOURCE_LABEL = "AB Taxes in Europe Database (TEDB)"
TEDB_BASE = "https://ec.europa.eu/taxation_customs/tedb/rest-api"
TEDB_UI_URL = "https://ec.europa.eu/taxation_customs/tedb/"
ALLOWED_HOSTS = ("ec.europa.eu",)
PARSER_VERSION = 1
# AB-27 tam liste; bundan az satır gelirse tohum korunur.
MIN_ROWS_FOR_REPLACE = 27
_MAX_BYTES = 12 * 1024 * 1024
_MAX_REDIRECTS = 5
DEFAULT_SYNC_INTERVAL = int(os.environ.get("EU_VAT_SYNC_INTERVAL_SECONDS", "86400"))

# Yunanistan AB belgelerinde ``EL``, ISO 3166-1'de ``GR``. TEDB ``EL`` kullanıyor;
# countries.py ve ihracat akışı ``GR`` kullanıyor. Eşleme iki yönlü olmak zorunda.
_ISO_ALIASES = {"EL": "GR", "GR": "EL"}

EU27 = frozenset({
    "AT", "BE", "BG", "HR", "CY", "CZ", "DK", "EE", "FI", "FR", "DE", "GR", "HU", "IE",
    "IT", "LV", "LT", "LU", "MT", "NL", "PL", "PT", "RO", "SK", "SI", "ES", "SE",
})

SUGGESTION_NOTE = (
    "Hedef ülkenin KDV oranı bilgilendirme amaçlıdır; beyannameyi açacak tarafın kendi "
    "mevzuatına göre doğrulaması gerekir."
)

# Uzman teyidi kaç gün geçerli sayılır? Oranlar değişir; süresiz "teyitli" demek bir süre
# sonra yalan olur. Bu yaştan sonra teyit düşer ve uyarı metni geri gelir.
CONFIRMATION_MAX_AGE_DAYS = int(os.environ.get("EU_VAT_CONFIRMATION_MAX_AGE_DAYS", "180"))


def _parse_day(value: Any) -> date | None:
    if not isinstance(value, str):
        return None
    try:
        return date.fromisoformat(value.strip()[:10])
    except ValueError:
        return None


def confirmation_state(
    payload: dict[str, Any],
    row: dict[str, Any] | None = None,
    *,
    today: date | None = None,
) -> dict[str, Any]:
    """Uzman teyidinin durumu: var mı, ne zaman verildi, kaç günlük, bayat mı.

    Tarih yoksa veya ayrıştırılamıyorsa teyit **yok** sayılır (güvenli taraf); ileri tarihli
    damga da bayat sayılır, çünkü güvenilmez.
    """
    source = row if row is not None else {}
    confirmed = bool(source.get("expert_confirmed", payload.get("expert_confirmed")))
    stamp = _parse_day(source.get("expert_confirmed_at") or payload.get("expert_confirmed_at"))
    if not confirmed or stamp is None:
        return {"confirmed": False, "confirmed_at": None, "age_days": None, "stale": False, "by": None}
    age = ((today or datetime.now(timezone.utc).date()) - stamp).days
    return {
        "confirmed": True,
        "confirmed_at": stamp.isoformat(),
        "age_days": age,
        "stale": age < 0 or age > CONFIRMATION_MAX_AGE_DAYS,
        "by": payload.get("expert_confirmed_by") or None,
    }


def _now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def _digits(value: Any) -> str:
    return "".join(ch for ch in str(value or "") if ch.isdigit())


def normalise_iso2(value: Any) -> str:
    """Ülke kodunu iki harfli büyük harfe indirger; ``EL`` girdisi ``GR``'ye çevrilir."""
    code = str(value or "").strip().upper()[:2]
    return "GR" if code == "EL" else code


def _default_data_dir() -> Path:
    override = os.environ.get("OFFICIAL_DATA_DIR")
    if override:
        return Path(override)
    packaged = Path(__file__).resolve().parent / "data" / "official"
    if packaged.is_dir():
        return packaged
    return Path.cwd() / "data" / "official"


def _default_cache_dir() -> Path:
    override = os.environ.get("MEVZUAT_DATA_DIR")
    if override:
        return Path(override)
    return Path.home() / ".cache" / "mevzuat-mcp"


# --- fasıl kuralı ---------------------------------------------------------------------


def _rule_matches(rule: dict[str, Any], gtip: str) -> bool:
    """Kural bu koda uyuyor mu? Fasıl aralığı veya pozisyon ön eki ile eşleşir."""
    code = _digits(gtip)
    if len(code) < 4:
        return False
    for heading in rule.get("headings") or []:
        if code.startswith(str(heading)):
            return True
    excluded = tuple(str(item) for item in rule.get("exclude_headings") or [])
    if excluded and code.startswith(excluded):
        return False
    chapter = int(code[:2])
    for pair in rule.get("chapters") or []:
        try:
            low, high = int(pair[0]), int(pair[1])
        except (TypeError, ValueError, IndexError):
            continue
        if low <= chapter <= high:
            return True
    return False


def match_chapter_rule(rules: list[dict[str, Any]], gtip: Any) -> dict[str, Any] | None:
    """Koda uyan ilk Ek-III kuralını döndürür; yoksa ``None``."""
    code = _digits(gtip)
    if len(code) < 4:
        return None
    for rule in rules:
        if _rule_matches(rule, code):
            return rule
    return None


# --- indeks ---------------------------------------------------------------------------


class EuVatRates:
    """AB-27 KDV oranları; tohumdan okur, TEDB çalışırsa resmî veriyle değiştirir."""

    def __init__(self, data_dir: str | Path | None = None, cache_dir: str | Path | None = None) -> None:
        self._seed_path = Path(data_dir or _default_data_dir()) / DATA_FILE
        self._cache_path = Path(cache_dir or _default_cache_dir()) / DATA_FILE
        self._payload: dict[str, Any] = {}
        self._rows: dict[str, dict[str, Any]] = {}
        self._rules: list[dict[str, Any]] = []
        self._origin = "none"
        self._last_error: str | None = None
        self._lock = asyncio.Lock()
        self.sync_interval = DEFAULT_SYNC_INTERVAL
        self._load()

    # -- yükleme

    def _read(self, path: Path) -> dict[str, Any] | None:
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            return None
        if not isinstance(payload, dict) or not isinstance(payload.get("rows"), list):
            return None
        return payload

    def _load(self) -> None:
        cached = self._read(self._cache_path)
        if (
            cached
            and cached.get("parser_version") == PARSER_VERSION
            and len(cached.get("rows") or []) >= MIN_ROWS_FOR_REPLACE
        ):
            self._install(cached, "synced")
            return
        seed = self._read(self._seed_path)
        if seed:
            self._install(seed, "seed")
            return
        self._payload, self._rows, self._rules, self._origin = {}, {}, [], "none"

    def _install(self, payload: dict[str, Any], origin: str) -> None:
        self._payload = payload
        self._origin = origin
        self._rules = [item for item in payload.get("chapter_rules") or [] if isinstance(item, dict)]
        rows: dict[str, dict[str, Any]] = {}
        for row in payload.get("rows") or []:
            if not isinstance(row, dict):
                continue
            iso2 = normalise_iso2(row.get("iso2"))
            if iso2 in EU27:
                rows[iso2] = row
        self._rows = rows

    @property
    def ready(self) -> bool:
        return bool(self._rows)

    def confirmation(self, row: dict[str, Any] | None = None) -> dict[str, Any]:
        """Bu satır (veya tablonun tamamı) için uzman teyidinin durumu."""
        return confirmation_state(self._payload, row)

    def _confirmed_fresh(self, row: dict[str, Any] | None = None) -> bool:
        state = self.confirmation(row)
        return bool(state["confirmed"]) and not state["stale"]

    def status(self) -> dict[str, Any]:
        # ``unverified`` anlamını korur: makine okuması yok demektir, teyit yok demek değil.
        unverified = [row["iso2"] for row in self._rows.values() if not row.get("verified")]
        confirmed = [
            row["iso2"] for row in self._rows.values() if self.confirmation(row)["confirmed"]
        ]
        # Oranı sık değişen ülke ancak teyit yoksa veya bayatsa uyarıya dönüşür.
        volatile = [
            row["iso2"]
            for row in self._rows.values()
            if row.get("volatile") and not self._confirmed_fresh(row)
        ]
        table = self.confirmation()
        return {
            "ready": self.ready,
            "row_count": len(self._rows),
            "source": self._payload.get("source") or SOURCE_LABEL,
            "source_url": self._payload.get("source_url") or TEDB_UI_URL,
            "retrieved_at": self._payload.get("retrieved_at"),
            "sha256": self._payload.get("sha256"),
            "parser_version": self._payload.get("parser_version"),
            "origin": self._origin,
            "checked_at": self._payload.get("checked_at"),
            "unverified": sorted(unverified),
            "expert_confirmed": sorted(confirmed),
            "expert_confirmed_at": table["confirmed_at"],
            "expert_confirmed_by": table["by"],
            "confirmation_age_days": table["age_days"],
            "confirmation_stale": table["stale"],
            "verify_first": sorted(volatile),
            "last_error": self._last_error,
            "cache_path": str(self._cache_path),
        }

    # -- sorgu

    def lookup(self, iso2: Any, *, gtip: Any = None) -> dict[str, Any]:
        """Bir üye devletin KDV oranı; GTİP verilirse Ek-III fasıl kuralı da uygulanır.

        Fasıl kuralı tetiklenirse ``applicable`` en düşük indirimli orana iner **ama**
        ``ambiguous`` her zaman ``True`` olur ve ``candidates`` standart oranı da taşır.
        Hangi oranın gerçekten uygulanacağını üye devletin kendi mevzuatı belirler.
        """
        code = normalise_iso2(iso2)
        row = self._rows.get(code)
        base = {
            "iso2": code,
            "country": None,
            "standard": None,
            "reduced": [],
            "super_reduced": None,
            "parking": None,
            "applicable": None,
            "applicable_basis": "unavailable",
            "candidates": [],
            "ambiguous": False,
            "matched_rule": None,
            "legal_basis": None,
            "verified": False,
            "verify_first": False,
            "expert_confirmed": False,
            "expert_confirmed_at": None,
            "confirmation_stale": False,
            "source": self._payload.get("source") or SOURCE_LABEL,
            "source_url": self._payload.get("source_url") or TEDB_UI_URL,
            "authority_url": None,
            "retrieved_at": self._payload.get("retrieved_at"),
            "origin": self._origin,
            "note": SUGGESTION_NOTE,
        }
        if row is None:
            base["note"] = (
                f"{code or 'Bu ülke'} AB üyesi değil veya KDV verimizde yok; oran gösterilmiyor."
                if code not in EU27
                else "Bu üye devlet için KDV satırı yüklenemedi."
            )
            return base

        state = self.confirmation(row)
        standard = row.get("standard")
        reduced = [float(item) for item in row.get("reduced") or []]
        base.update(
            {
                "country": row.get("country"),
                "standard": standard,
                "reduced": reduced,
                "super_reduced": row.get("super_reduced"),
                "parking": row.get("parking"),
                "applicable": standard,
                "applicable_basis": "standard",
                "candidates": [standard] if standard is not None else [],
                "verified": bool(row.get("verified")),
                "verify_first": bool(row.get("volatile")) and not self._confirmed_fresh(row),
                "expert_confirmed": state["confirmed"],
                "expert_confirmed_at": state["confirmed_at"],
                "confirmation_stale": state["stale"],
                "authority_url": row.get("authority_url"),
            }
        )

        rule = match_chapter_rule(self._rules, gtip) if gtip else None
        if rule and reduced:
            lowest = min(reduced)
            base.update(
                {
                    "applicable": lowest,
                    "applicable_basis": "chapter_rule",
                    "ambiguous": True,
                    "matched_rule": rule.get("label") or rule.get("id"),
                    "legal_basis": rule.get("legal_basis"),
                    "candidates": sorted({*reduced, *( [standard] if standard is not None else [] )}),
                    "note": (
                        f"{rule.get('label')} kalemi AB Ek-III kapsamında indirimli orana uygun olabilir. "
                        f"Hangi oranın uygulanacağını {row.get('country')} mevzuatı belirler; "
                        f"standart oran %{standard} da geçerli olabilir. "
                        + ("Bu kalem Ek-III'te dar tanımlıdır, kapsam dışı kalabilir. " if rule.get("narrow") else "")
                        + SUGGESTION_NOTE
                    ),
                }
            )
        elif rule and not reduced:
            base["note"] = (
                f"{row.get('country')} indirimli oran uygulamıyor; {rule.get('label')} kaleminde de "
                f"standart oran geçerlidir. " + SUGGESTION_NOTE
            )
        authority = row.get("authority_url") or "resmî vergi idaresi"
        if base["verified"]:
            pass  # TEDB'den makine okuması; ek köken uyarısı gerekmez.
        elif state["confirmed"] and not state["stale"]:
            base["note"] = (
                f"Bu oran {state['confirmed_at']} tarihinde uzman teyidiyle güncel kabul edildi"
                + (f" ({state['by']})" if state["by"] else "")
                + f"; {authority} üzerinden doğrulayabilirsiniz. "
                + base["note"]
            )
        else:
            base["note"] = (
                "Bu oran doğrulanmamış tohum veriden geliyor"
                + (" ve bu ülkenin oranı son iki yılda değişti" if row.get("volatile") else "")
                + (
                    f"; uzman teyidi {state['age_days']} gün önce yapıldı, tazelenmeli"
                    if state["confirmed"]
                    else ""
                )
                + f"; {authority} üzerinden teyit edin. "
                + base["note"]
            )
        return base

    def summary_lines(self, report: dict[str, Any]) -> list[str]:
        """Sonucu insan okunur birkaç satıra indirger (``vat_lists.summary_lines`` deseni)."""
        if report.get("standard") is None:
            return [report.get("note") or "Hedef ülke KDV verisi yok."]
        lines = [f"{report.get('country')} standart KDV oranı: %{report['standard']}"]
        if report.get("applicable_basis") == "chapter_rule":
            lines.append(
                f"{report.get('matched_rule')} kuralı nedeniyle indirimli %{report['applicable']} "
                f"uygulanabilir (adaylar: {', '.join('%' + str(item) for item in report.get('candidates') or [])})"
            )
        if report.get("legal_basis"):
            lines.append(str(report["legal_basis"]))
        lines.append(str(report.get("note") or ""))
        return [line for line in lines if line]

    # -- eşitleme

    def _own_client(self) -> httpx.AsyncClient:
        return httpx.AsyncClient(
            timeout=httpx.Timeout(90, connect=20),
            headers={"User-Agent": USER_AGENT, "Accept-Language": "en"},
            verify=official_ssl_context(),
            follow_redirects=False,
        )

    async def _get(self, client: httpx.AsyncClient, url: str) -> bytes:
        """Her yönlendirme adımında egress kontrolü yeniden yapılır (foreign_tariff deseni)."""
        current = url
        for _ in range(_MAX_REDIRECTS):
            validate_outbound_url(current, allowed_hosts=ALLOWED_HOSTS)
            response = await client.get(current)
            if response.is_redirect:
                location = response.headers.get("location", "")
                if not location:
                    raise ValueError("TEDB hedefsiz yönlendirme döndürdü.")
                current = httpx.URL(str(response.url)).join(location).human_repr()
                continue
            response.raise_for_status()
            content = response.content
            if len(content) > _MAX_BYTES:
                raise ValueError("TEDB yanıtı beklenenden büyük.")
            return content
        raise ValueError("TEDB çok fazla yönlendirme yaptı.")

    def parse_export(self, payload: bytes) -> list[dict[str, Any]]:
        """TEDB dışa aktarım çalışma kitabını ülke başına orana çevirir.

        Boş şablon (bugünkü durum) sıfır satır döndürür; çağıran tohumu korur.
        """
        import openpyxl

        if not payload.startswith(b"PK\x03\x04"):
            raise ValueError("TEDB yanıtı bir çalışma kitabı değil.")
        workbook = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        try:
            sheet = workbook["Results"] if "Results" in workbook.sheetnames else workbook.worksheets[0]
            collected: dict[str, dict[str, Any]] = {}
            header_seen = False
            for raw in sheet.iter_rows(values_only=True):
                values = [("" if cell is None else str(cell).strip()) for cell in raw]
                values += [""] * (11 - len(values))
                if not header_seen:
                    header_seen = values[0].lower() == "country"
                    continue
                iso2 = normalise_iso2(values[0])
                if iso2 not in EU27:
                    continue
                rate_type = values[2].strip().lower()
                try:
                    rate = float(values[3].replace(",", "."))
                except (TypeError, ValueError):
                    continue
                entry = collected.setdefault(
                    iso2,
                    {"iso2": iso2, "tedb_code": _ISO_ALIASES.get(iso2, iso2), "standard": None,
                     "reduced": [], "super_reduced": None, "parking": None, "verified": True,
                     "volatile": False},
                )
                if "standard" in rate_type:
                    entry["standard"] = rate
                elif "super" in rate_type:
                    entry["super_reduced"] = rate
                elif "parking" in rate_type:
                    entry["parking"] = rate
                elif "reduced" in rate_type and rate not in entry["reduced"]:
                    entry["reduced"].append(rate)
            return [row for row in collected.values() if row["standard"] is not None]
        finally:
            workbook.close()

    async def sync(self, http: httpx.AsyncClient | None = None) -> dict[str, Any]:
        """TEDB'yi yoklar; veri gelirse tohumu resmî veriyle değiştirir. Asla hata fırlatmaz."""
        async with self._lock:
            client = http or self._own_client()
            try:
                config = json.loads(await self._get(client, f"{TEDB_BASE}/configurations"))
                ids = [
                    str(country.get("id"))
                    for country in config.get("countries") or []
                    if normalise_iso2(country.get("defaultCountryCode")) in EU27
                ]
                if not ids:
                    raise ValueError("TEDB ülke listesi boş döndü.")
                today = datetime.now(timezone.utc).date().isoformat()
                query = urlencode({"dateFrom": today, "dateTo": today, "memberStates": ",".join(ids)})
                content = await self._get(client, f"{TEDB_BASE}/vatSearch/export?{query}")
                rows = self.parse_export(content)
            except (SecurityViolation, httpx.HTTPError, ValueError, KeyError, TypeError) as exc:
                message = f"{type(exc).__name__}: {str(exc)[:160]}"
                self._last_error = message
                logger.warning("AB KDV oranları eşitlenemedi; tohum korunuyor: %s", message)
                return {"ok": False, "error": message, "row_count": len(self._rows)}
            finally:
                if http is None:
                    await client.aclose()

            if len(rows) < MIN_ROWS_FOR_REPLACE:
                # Bugünkü durum: TEDB boş şablon veriyor. Tohumu asla bununla ezmeyiz.
                message = (
                    f"TEDB yalnız {len(rows)} satır döndürdü (en az {MIN_ROWS_FOR_REPLACE} gerekir); "
                    "tohum korunuyor."
                )
                self._last_error = message
                logger.info("AB KDV eşitlemesi atlandı: %s", message)
                return {"ok": False, "error": message, "row_count": len(self._rows)}

            payload = {
                "source": SOURCE_LABEL,
                "source_url": TEDB_UI_URL,
                "parser_version": PARSER_VERSION,
                "retrieved_at": _now(),
                "sha256": hashlib.sha256(content).hexdigest(),
                "note": "TEDB dışa aktarımından otomatik ayrıştırıldı.",
                "checked_at": datetime.now(timezone.utc).date().isoformat(),
                "chapter_rules": self._rules,
                "rows": rows,
            }
            self._install(payload, "synced")
            self._last_error = None
            try:
                self._cache_path.parent.mkdir(parents=True, exist_ok=True)
                tmp = self._cache_path.with_suffix(".tmp")
                tmp.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
                tmp.replace(self._cache_path)
            except OSError as exc:
                logger.warning("AB KDV önbelleği yazılamadı: %s", exc)
            return {
                "ok": True,
                "row_count": len(rows),
                "sha256": payload["sha256"],
                "retrieved_at": payload["retrieved_at"],
            }

    async def periodic_sync_loop(self, *, initial_delay: float = 120.0) -> None:
        if self.sync_interval <= 0:
            return
        await asyncio.sleep(initial_delay)
        while True:
            try:
                await self.sync()
            except Exception:  # döngü hiçbir koşulda ölmemeli
                logger.exception("AB KDV eşitleme döngüsü hata verdi")
            await asyncio.sleep(max(3600, self.sync_interval))


__all__ = [
    "EU27",
    "EuVatRates",
    "match_chapter_rule",
    "normalise_iso2",
]
